from __future__ import annotations

import copy
import hashlib
import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.models import ListingProject, ListingProjectInput, ListingVersion


TEMPLATE_PATH = Path(__file__).parent / "templates" / "amazon" / "SAW_BLADE.xlsm"
TEMPLATE_SHEET = "Template"
DATA_START_ROW = 8


@dataclass
class ExportedAmazonTemplate:
    file_name: str
    content: bytes


def build_saw_blade_upload_workbook(
    *,
    project: ListingProject,
    inputs: ListingProjectInput | None,
    listing: ListingVersion,
    brand: str,
    category: str,
    images: list[dict[str, Any]],
    variation: dict[str, Any] | None = None,
) -> ExportedAmazonTemplate:
    if not TEMPLATE_PATH.exists():
        raise RuntimeError("SAW_BLADE 模板不存在，请确认 app/templates/amazon/SAW_BLADE.xlsm 已部署。")

    workbook = load_workbook(TEMPLATE_PATH, keep_vba=True)
    sheet = workbook[TEMPLATE_SHEET]
    columns = _template_columns(sheet)

    rows = _build_rows(
        project=project,
        inputs=inputs,
        listing=listing,
        brand=brand,
        category=category,
        images=images,
        variation=variation,
    )

    _clear_existing_data(sheet)
    for offset, row_data in enumerate(rows):
        row_number = DATA_START_ROW + offset
        _copy_style_from_template_row(sheet, row_number)
        for field_name, value in row_data.items():
            column = columns.get(field_name)
            if not column:
                continue
            sheet.cell(row=row_number, column=column).value = value
    _clear_suppressed_columns(sheet, columns, len(rows))

    stream = BytesIO()
    workbook.save(stream)
    file_name = f"{_slug(rows[0].get('contribution_sku#1.value') or project.project_name)}-SAW_BLADE-amazon-upload.xlsm"
    return ExportedAmazonTemplate(file_name=file_name, content=stream.getvalue())


def _template_columns(sheet) -> dict[str, int]:
    columns: dict[str, int] = {}
    for cell in sheet[5]:
        if cell.value:
            columns[str(cell.value).strip()] = cell.column
    return columns


def _clear_existing_data(sheet) -> None:
    if sheet.max_row > DATA_START_ROW:
        sheet.delete_rows(DATA_START_ROW + 1, sheet.max_row - DATA_START_ROW)
    for cell in sheet[DATA_START_ROW]:
        cell.value = None


def _copy_style_from_template_row(sheet, row_number: int) -> None:
    if row_number == DATA_START_ROW:
        return
    for column in range(1, sheet.max_column + 1):
        source = sheet.cell(row=DATA_START_ROW, column=column)
        target = sheet.cell(row=row_number, column=column)
        if source.has_style:
            target._style = copy.copy(source._style)
        target.font = copy.copy(source.font)
        target.fill = copy.copy(source.fill)
        target.border = copy.copy(source.border)
        target.alignment = copy.copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy.copy(source.protection)


def _clear_suppressed_columns(sheet, columns: dict[str, int], row_count: int) -> None:
    suppressed_fields = [
        "item_type_keyword[marketplace_id=ATVPDKIKX0DER]#1.value",
    ]
    for field_name in suppressed_fields:
        column = columns.get(field_name)
        if not column:
            continue
        for row_number in range(DATA_START_ROW, DATA_START_ROW + row_count):
            sheet.cell(row=row_number, column=column).value = None


def _build_rows(
    *,
    project: ListingProject,
    inputs: ListingProjectInput | None,
    listing: ListingVersion,
    brand: str,
    category: str,
    images: list[dict[str, Any]],
    variation: dict[str, Any] | None,
) -> list[dict[str, Any]]:
    base = _base_row(project, inputs, listing, brand, category, images)
    if not variation or not variation.get("enabled"):
        return [base]

    clean_variants = [
        item for item in variation.get("variants") or []
        if any(_text(item.get(key)) for key in ("sku", "value", "color", "price"))
    ]
    if not clean_variants:
        return [base]

    theme = _amazon_variation_theme(_text(variation.get("theme")) or "SIZE")
    parent_sku = _explicit_sku(_text(variation.get("parentSku"))) or _parent_sku(project)
    parent_row = {
        "contribution_sku#1.value": parent_sku,
        "product_type#1.value": "SAW_BLADE",
        "::record_action": "Create or Replace (Full Update)",
        "parentage_level[marketplace_id=ATVPDKIKX0DER]#1.value": "Parent",
        "variation_theme#1.name": theme,
        "item_name[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": base["item_name[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
        "brand[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": brand,
        "package_level[marketplace_id=ATVPDKIKX0DER]#1.value": "Unit",
        "manufacturer[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": brand,
        "country_of_origin[marketplace_id=ATVPDKIKX0DER]#1.value": "China",
    }
    child_rows = []
    for index, variant in enumerate(clean_variants, start=1):
        value = _text(variant.get("value") or variant.get("color") or f"Variant {index}")
        child = dict(base)
        child["contribution_sku#1.value"] = _explicit_sku(_text(variant.get("sku"))) or _variant_sku(parent_sku, value, index)
        child["parentage_level[marketplace_id=ATVPDKIKX0DER]#1.value"] = "Child"
        child["child_parent_sku_relationship[marketplace_id=ATVPDKIKX0DER]#1.parent_sku"] = parent_sku
        child["variation_theme#1.name"] = theme
        child["item_name[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"] = _join_title(base["item_name[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"], value)
        if variant.get("price"):
            child["list_price[marketplace_id=ATVPDKIKX0DER]#1.value"] = _number(variant.get("price"))
            child["purchasable_offer[marketplace_id=ATVPDKIKX0DER][audience=ALL]#1.our_price#1.schedule#1.value_with_tax"] = _number(variant.get("price"))
        if "SIZE" in theme:
            child["size[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"] = value
        if "COLOR" in theme:
            child["color[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"] = _text(variant.get("color") or value)
        if "NUMBER_OF_ITEMS" in theme:
            qty = _digits(value) or value
            child["number_of_items[marketplace_id=ATVPDKIKX0DER]#1.value"] = qty
            child["item_package_quantity[marketplace_id=ATVPDKIKX0DER]#1.value"] = qty
        _apply_variant_image(child, images, _text(variant.get("imageType")))
        child_rows.append(child)
    return [parent_row, *child_rows]


def _base_row(
    project: ListingProject,
    inputs: ListingProjectInput | None,
    listing: ListingVersion,
    brand: str,
    category: str,
    images: list[dict[str, Any]],
) -> dict[str, Any]:
    bullets = [listing.bullet_1, listing.bullet_2, listing.bullet_3, listing.bullet_4, listing.bullet_5]
    row: dict[str, Any] = {
        "contribution_sku#1.value": _single_sku(project, listing),
        "product_type#1.value": "SAW_BLADE",
        "::record_action": "Create or Replace (Full Update)",
        "item_name[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": _limit(listing.title or project.product_name or project.project_name, 200),
        "brand[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": brand,
        "amzn1.volt.ca.product_id_type": "GTIN Exempt",
        "package_level[marketplace_id=ATVPDKIKX0DER]#1.value": "Unit",
        "manufacturer[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": brand,
        "product_description[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": _limit(listing.description, 2000),
        "generic_keyword[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": _limit(listing.backend_search_terms or _keywords_from_inputs(inputs), 250),
        "condition_type[marketplace_id=ATVPDKIKX0DER]#1.value": "New",
        "country_of_origin[marketplace_id=ATVPDKIKX0DER]#1.value": "China",
        "list_price[marketplace_id=ATVPDKIKX0DER]#1.value": _number(project.target_price),
        "purchasable_offer[marketplace_id=ATVPDKIKX0DER][audience=ALL]#1.our_price#1.schedule#1.value_with_tax": _number(project.target_price),
        "size[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": _text(getattr(inputs, "size", "")),
        "material[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value": _text(getattr(inputs, "material", "")),
        "number_of_items[marketplace_id=ATVPDKIKX0DER]#1.value": _digits(_text(getattr(inputs, "quantity", ""))),
        "item_package_quantity[marketplace_id=ATVPDKIKX0DER]#1.value": _digits(_text(getattr(inputs, "quantity", ""))),
    }
    for index, bullet in enumerate(bullets, start=1):
        row[f"bullet_point[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#{index}.value"] = _limit(bullet, 500)
    _apply_images(row, images)
    return row


def _apply_variant_image(row: dict[str, Any], images: list[dict[str, Any]], image_type: str) -> None:
    selected = next((item for item in images if item.get("image_type") == image_type and item.get("url")), None)
    if not selected:
        return
    reordered = [selected, *[item for item in images if item.get("id") != selected.get("id")]]
    _apply_images(row, reordered)


def _apply_images(row: dict[str, Any], images: list[dict[str, Any]]) -> None:
    ordered = _ordered_images(images)[:9]
    if ordered:
        row["main_product_image_locator[marketplace_id=ATVPDKIKX0DER]#1.media_location"] = ordered[0]["url"]
    for index, image in enumerate(ordered[1:9], start=1):
        row[f"other_product_image_locator_{index}[marketplace_id=ATVPDKIKX0DER]#1.media_location"] = image["url"]


def _ordered_images(images: list[dict[str, Any]]) -> list[dict[str, Any]]:
    priority = {
        "main_image": 0,
        "dimension": 1,
        "feature": 2,
        "compatibility": 3,
        "application": 4,
        "package": 5,
        "other": 6,
        "aplus": 99,
    }
    return sorted(
        [item for item in images if item.get("url")],
        key=lambda item: (priority.get(item.get("image_type"), 50), item.get("id") or 0),
    )


def _amazon_variation_theme(value: str) -> str:
    mapping = {
        "SizeName": "SIZE",
        "ColorName": "COLOR",
        "SizeName-ColorName": "COLOR/SIZE",
        "PackageQuantity": "NUMBER_OF_ITEMS",
    }
    return mapping.get(value, value or "SIZE")


def _keywords_from_inputs(inputs: ListingProjectInput | None) -> str:
    parts = [
        getattr(inputs, "main_keywords", ""),
        getattr(inputs, "secondary_keywords", ""),
        getattr(inputs, "long_tail_keywords", ""),
        getattr(inputs, "compatibility_keywords", ""),
    ]
    return " ".join(_text(item) for item in parts if _text(item))


def _single_sku(project: ListingProject, listing: ListingVersion) -> str:
    seed = f"project:{getattr(project, 'id', '')}|listing:{getattr(listing, 'id', '')}|{project.project_name}|{project.product_name}"
    return _hashed_sku(project.product_name or project.project_name, seed=seed)


def _parent_sku(project: ListingProject) -> str:
    seed = f"project:{getattr(project, 'id', '')}|parent|{project.project_name}|{project.product_name}"
    return _hashed_sku(project.product_name or project.project_name or "PARENT", seed=seed, suffix="P")


def _variant_sku(parent_sku: str, value: str, index: int) -> str:
    value_slug = _slug(value)
    digest = _hash(f"{parent_sku}|{value}|{index}")[:6]
    parts = [parent_sku[:24], value_slug[:8] or f"V{index}", digest]
    return "-".join(part for part in parts if part)[:40]


def _explicit_sku(value: str) -> str:
    if not _text(value):
        return ""
    text = _slug(value).upper()
    return text[:40] if text else ""


def _hashed_sku(value: Any, *, seed: str, suffix: str = "") -> str:
    base = _slug(value).upper()
    if not base or base == "AMAZON-LISTING":
        base = "SKU"
    digest = _hash(seed or value)[:8]
    pieces = [base[:26], suffix, digest]
    return "-".join(piece for piece in pieces if piece)[:40]


def _slug(value: Any) -> str:
    text = _text(value)
    text = re.sub(r"[^A-Za-z0-9]+", "-", text).strip("-")
    return text or "AMAZON-LISTING"


def _hash(value: Any) -> str:
    return hashlib.sha1(_text(value).encode("utf-8")).hexdigest().upper()


def _join_title(title: Any, suffix: Any) -> str:
    suffix_text = _text(suffix)
    if not suffix_text or suffix_text.lower() in _text(title).lower():
        return _limit(title, 200)
    return _limit(f"{title} - {suffix_text}", 200)


def _digits(value: str) -> str:
    match = re.search(r"\d+(?:\.\d+)?", value)
    return match.group(0) if match else ""


def _number(value: Any) -> Any:
    if value in (None, ""):
        return None
    try:
        return round(float(str(value).replace("$", "").strip()), 2)
    except ValueError:
        return value


def _limit(value: Any, limit: int) -> str:
    return _text(value)[:limit]


def _text(value: Any) -> str:
    return str(value or "").replace("\r", " ").replace("\n", " ").strip()
